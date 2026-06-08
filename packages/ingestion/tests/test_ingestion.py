"""Unit tests for ingestion.

Each test maps to a scenario in
openspec/changes/add-ingestion/specs/ingestion/spec.md.
"""

from __future__ import annotations

import csv
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from ingestion import ingest, ingest_csv, ingest_path, ingest_xlsx


def _write_csv(path: Path, rows):
    with path.open("w", newline="", encoding="utf-8") as handle:
        csv.writer(handle).writerows(rows)


def _write_xlsx(path: Path, sheets: dict[str, list[list[str]]]):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    wb.save(str(path))


def test_csv_to_table(tmp_path):
    p = tmp_path / "d.csv"
    _write_csv(p, [["Sample", "87Sr/86Sr"], ["HT-1", "0.7041"], ["HT-2", "0.7039"]])
    res = ingest_csv(p)
    assert len(res.tables) == 1
    table = res.tables[0]
    assert table.columns == ["Sample", "87Sr/86Sr"]
    assert table.rows == [["HT-1", "0.7041"], ["HT-2", "0.7039"]]
    assert res.assets[0].kind == "table"
    assert res.assets[0].source == "d.csv"


def test_xlsx_two_sheets(tmp_path):
    p = tmp_path / "d.xlsx"
    _write_xlsx(p, {"S1": [["a", "b"], ["1", "2"]], "S2": [["x"], ["9"]]})
    res = ingest_xlsx(p)
    assert len(res.tables) == 2
    assert {a.locator.get("sheet") for a in res.assets} == {"S1", "S2"}


def test_provenance_records_source_and_locator(tmp_path):
    p = tmp_path / "d.xlsx"
    _write_xlsx(p, {"Data": [["h"], ["1"]]})
    asset = ingest_xlsx(p).assets[0]
    assert asset.source == "d.xlsx"
    assert asset.locator == {"sheet": "Data"}


# --- add-parse-cache ---------------------------------------------------------


def test_parse_cache_hits_on_second_call(tmp_path):
    from ingestion.cache import cached_pdf
    from ingestion.models import IngestResult
    from slide_ir import EvidenceAsset

    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4 fake bytes")
    calls = {"n": 0}

    def parse(_p):
        calls["n"] += 1
        return IngestResult(
            assets=[EvidenceAsset(asset_id="x", kind="section_text", content_ref="hello", source="a.pdf", locator={"page": 1})]
        )

    cache = tmp_path / "cache"
    r1 = cached_pdf(pdf, parser_key="auto", cache_dir=cache, parse_fn=parse)
    r2 = cached_pdf(pdf, parser_key="auto", cache_dir=cache, parse_fn=parse)
    assert calls["n"] == 1  # second call was a cache hit
    assert r1.assets[0].content_ref == "hello" and r2.assets[0].content_ref == "hello"


# --- add-parser-resilience: quality gating + cascade -------------------------


def test_assess_quality_flags_thin_parse():
    from ingestion import assess_quality
    from ingestion.models import IngestResult
    from slide_ir import EvidenceAsset

    thin = IngestResult(
        assets=[EvidenceAsset(asset_id="p1", kind="section_text", content_ref="short", source="x.pdf", locator={"page": 1})]
    )
    q = assess_quality(thin)
    assert q["adequate"] is False and q["warnings"]

    good = IngestResult(
        assets=[EvidenceAsset(asset_id="p1", kind="section_text", content_ref="x" * 1000, source="x.pdf", locator={"page": 1})]
    )
    assert assess_quality(good)["adequate"] is True


def test_cascade_descends_on_thin_parse(tmp_path, monkeypatch):
    import ingestion.router as router
    from ingestion.models import IngestResult
    from slide_ir import EvidenceAsset

    def thin(_p):
        return IngestResult(assets=[EvidenceAsset(asset_id="t", kind="section_text", content_ref="tiny", source="x.pdf", locator={"page": 1})])

    def good(_p):
        return IngestResult(assets=[EvidenceAsset(asset_id="g", kind="section_text", content_ref="y" * 1000, source="x.pdf", locator={"page": 1})])

    monkeypatch.setattr(router, "_pdf_backends", lambda forced, ws: [("mineru", thin), ("pdfplumber", good)])
    res = router._ingest_pdf(tmp_path / "x.pdf", tmp_path)
    assert any(len(a.content_ref or "") >= 1000 for a in res.assets)  # took the good (adequate) one
    assert any("降级" in w for w in res.warnings)


def test_docling_skipped_when_unavailable(monkeypatch):
    import ingestion.router as router
    from ingestion import docling_available

    monkeypatch.delenv("MINERU_API_KEY", raising=False)
    names = [n for n, _ in router._pdf_backends(None, "/tmp/ws")]
    assert "pdfplumber" in names
    if not docling_available():  # optional plugin: absent -> not in the cascade
        assert "docling" not in names


def test_zip_forwards_workspace(tmp_path, monkeypatch):
    import ingestion.router as router

    seen: list = []
    real = router.ingest_path

    def spy(path, *, workspace=None, cache_dir=None):
        seen.append(workspace)
        return real(path, workspace=workspace, cache_dir=cache_dir)

    monkeypatch.setattr(router, "ingest_path", spy)
    csvp = tmp_path / "a.csv"
    _write_csv(csvp, [["h"], ["1"]])
    zp = tmp_path / "z.zip"
    with zipfile.ZipFile(zp, "w") as zf:
        zf.write(csvp, "a.csv")
    from ingestion import ingest_zip

    ws = tmp_path / "ws"
    ingest_zip(zp, workspace=ws)
    assert any(str(w) == str(ws) for w in seen if w is not None)  # member got the workspace


def test_zip_recurses(tmp_path):
    csvp = tmp_path / "a.csv"
    _write_csv(csvp, [["h"], ["1"]])
    xlp = tmp_path / "b.xlsx"
    _write_xlsx(xlp, {"S": [["h2"], ["2"]]})
    zp = tmp_path / "arch.zip"
    with zipfile.ZipFile(zp, "w") as zf:
        zf.write(csvp, "a.csv")
        zf.write(xlp, "b.xlsx")
    res = ingest_path(zp)
    assert len(res.tables) == 2
    assert all(a.source.startswith("arch.zip!") for a in res.assets)


def test_unknown_skipped_and_multi_combine(tmp_path):
    txt = tmp_path / "n.txt"
    txt.write_text("hello", encoding="utf-8")
    assert ingest_path(txt).tables == []

    c1 = tmp_path / "a.csv"
    _write_csv(c1, [["h"], ["1"]])
    c2 = tmp_path / "b.csv"
    _write_csv(c2, [["h"], ["2"]])
    res = ingest(c1, c2)
    assert len(res.tables) == 2
    assert [a.content_ref for a in res.assets] == ["table:0", "table:1"]  # re-based refs


def test_image_to_figure_asset(tmp_path):
    img = tmp_path / "f.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n")
    res = ingest_path(img)
    assert len(res.assets) == 1
    assert res.assets[0].kind == "figure"


def test_pdf_text_extraction(tmp_path):
    mpl = pytest.importorskip("matplotlib")
    mpl.use("Agg")
    import matplotlib.pyplot as plt

    pdf = tmp_path / "doc.pdf"
    fig = plt.figure()
    fig.text(0.1, 0.5, "HelloIngestion")
    fig.savefig(pdf)
    plt.close(fig)

    res = ingest_path(pdf)
    texts = [a.content_ref for a in res.assets if a.kind == "section_text"]
    assert any("HelloIngestion" in t for t in texts)


# --- improve-output-quality: junk-table filtering ----------------------------


def test_low_quality_table_is_filtered():
    from ingestion.models import is_low_quality_table
    from slide_ir import TableBlock

    # majority auto-named headers (pdfplumber noise) -> dropped
    assert is_low_quality_table(TableBlock(columns=["col1", "Se-", "col3"], rows=[["a", "b", "c"]]))
    assert is_low_quality_table(TableBlock(columns=["A", "B"], rows=[]))  # no data rows
    assert is_low_quality_table(TableBlock(columns=["A"], rows=[["x"]]))  # < 2 columns
    # a clean table is kept
    assert not is_low_quality_table(
        TableBlock(columns=["Sample", "87Sr/86Sr"], rows=[["HT-1", "0.7041"]])
    )


# --- add-figure-extraction ---------------------------------------------------


def _figure_pdf(path):
    mpl = pytest.importorskip("matplotlib")
    mpl.use("Agg")
    import matplotlib.pyplot as plt

    fig = plt.figure(figsize=(6, 7))
    fig.text(0.12, 0.85, "Introductory body text long enough to act as the prose boundary above figure.")
    ax = fig.add_axes([0.15, 0.4, 0.6, 0.35])
    ax.plot([0, 1, 2], [1, 3, 2])
    fig.text(0.12, 0.2, "Fig. 1. Synthetic test figure caption used for the extraction unit test here.")
    fig.savefig(path)
    plt.close(fig)


def test_figure_extraction_from_caption(tmp_path):
    pdf = tmp_path / "fig.pdf"
    _figure_pdf(pdf)
    from ingestion import extract_figures

    assets = extract_figures(pdf, tmp_path / "ws")
    assert assets, "expected at least one figure asset"
    a = assets[0]
    assert a.kind == "figure"
    assert Path(a.content_ref).is_file() and Path(a.content_ref).stat().st_size > 0
    assert "Fig. 1" in a.locator["caption"] and a.locator["page"] == 1


def test_no_caption_yields_no_figures(tmp_path):
    mpl = pytest.importorskip("matplotlib")
    mpl.use("Agg")
    import matplotlib.pyplot as plt

    pdf = tmp_path / "plain.pdf"
    fig = plt.figure()
    fig.text(0.1, 0.5, "JustBodyTextNoCaptions")
    fig.savefig(pdf)
    plt.close(fig)
    from ingestion import extract_figures

    assert extract_figures(pdf, tmp_path / "ws2") == []


# --- add-mineru-parser: content_list -> Evidence Pool (no network) ------------


def test_parse_mineru_content_list(tmp_path):
    from ingestion import parse_mineru_content_list

    assets_dir = tmp_path / "mineru"
    (assets_dir / "images").mkdir(parents=True)
    img = assets_dir / "images" / "fig1.jpg"
    img.write_bytes(b"\xff\xd8\xff\xe0JFIF")  # minimal jpeg-ish bytes

    blocks = [
        {"type": "text", "text": "Introduction", "text_level": 1, "page_idx": 0},
        {"type": "text", "text": "Atmospheric oxygen rose over billions of years.", "page_idx": 0},
        {"type": "equation", "text": "O_2 = f(t)", "text_format": "latex", "page_idx": 0},
        {
            "type": "table",
            "table_body": "<table><tr><td>Sample</td><td>Value</td></tr><tr><td>HT-1</td><td>0.7</td></tr></table>",
            "table_caption": ["Table 1. measurements"],
            "page_idx": 1,
        },
        # MinerU renders figures as type "chart"; caption may be empty -> borrowed from a "Fig. N" text
        {"type": "text", "text": "Fig. 1. Temporal trends of trace elements.", "page_idx": 1},
        {"type": "chart", "img_path": "images/fig1.jpg", "chart_caption": [], "page_idx": 1},
    ]
    res = parse_mineru_content_list(blocks, assets_dir, "paper.pdf", tmp_path / "ws", "paper")

    texts = [a for a in res.assets if a.kind == "section_text"]
    figs = [a for a in res.assets if a.kind == "figure"]
    assert len(res.tables) == 1 and res.tables[0].columns == ["Sample", "Value"]
    assert any("# Introduction" in a.content_ref and "[formula] O_2" in a.content_ref for a in texts)
    assert len(figs) == 1
    assert Path(figs[0].content_ref).is_file()
    assert "Fig. 1" in figs[0].locator["caption"]


# --- figure panel splitting (ASA_SPLIT_FIGURES) ------------------------------


def _panel_image(path, w, h, cols, rows, gutter=30):
    """Make a cols x rows grid of solid-color cells separated by white gutters."""
    from PIL import Image, ImageDraw

    im = Image.new("RGB", (w, h), (255, 255, 255))
    d = ImageDraw.Draw(im)
    cw = (w - (cols + 1) * gutter) // cols
    ch = (h - (rows + 1) * gutter) // rows
    colors = [(200, 60, 60), (60, 120, 200), (60, 170, 90), (180, 140, 40), (140, 70, 160), (70, 160, 160)]
    k = 0
    for r in range(rows):
        for c in range(cols):
            x = gutter + c * (cw + gutter)
            y = gutter + r * (ch + gutter)
            d.rectangle([x, y, x + cw, y + ch], fill=colors[k % len(colors)])
            k += 1
    im.save(str(path))
    return path


def test_split_composite_2x2(tmp_path):
    from ingestion import split_composite

    img = _panel_image(tmp_path / "grid.png", 800, 600, 2, 2)
    panels = split_composite(img, tmp_path / "ws", "fig")
    assert len(panels) == 4 and all(p.exists() for p in panels)


def test_split_composite_1x3(tmp_path):
    from ingestion import split_composite

    img = _panel_image(tmp_path / "strip.png", 900, 500, 3, 1)
    assert len(split_composite(img, tmp_path / "ws", "fig")) == 3


def test_split_composite_single_panel_no_split(tmp_path):
    from PIL import Image
    from ingestion import split_composite

    Image.new("RGB", (800, 600), (120, 120, 120)).save(str(tmp_path / "solid.png"))  # no gutters
    assert split_composite(tmp_path / "solid.png", tmp_path / "ws", "fig") == []


def test_split_composite_small_image_gated(tmp_path):
    from ingestion import split_composite

    img = _panel_image(tmp_path / "tiny.png", 200, 150, 2, 1)  # below MIN_W/MIN_H
    assert split_composite(img, tmp_path / "ws", "fig") == []


def test_mineru_panel_split_opt_in(tmp_path, monkeypatch):
    from ingestion import parse_mineru_content_list

    assets_dir = tmp_path / "mineru"
    (assets_dir / "images").mkdir(parents=True)
    _panel_image(assets_dir / "images" / "fig.png", 800, 600, 2, 2)
    blocks = [{"type": "chart", "img_path": "images/fig.png", "chart_caption": ["Fig. 1. grid"], "page_idx": 0}]

    # opt-in OFF -> only the whole figure
    off = parse_mineru_content_list(blocks, assets_dir, "p.pdf", tmp_path / "ws_off", "p")
    assert sum(1 for a in off.assets if a.kind == "figure") == 1

    # opt-in ON -> whole figure + 4 panels
    monkeypatch.setenv("ASA_SPLIT_FIGURES", "1")
    on = parse_mineru_content_list(blocks, assets_dir, "p.pdf", tmp_path / "ws_on", "p")
    figs = [a for a in on.assets if a.kind == "figure"]
    assert len(figs) == 5 and any(a.locator.get("panel") == 0 for a in figs)
